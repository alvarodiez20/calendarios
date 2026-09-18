"""Colores y estilos compartidos por los Excel de entrada y salida."""

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

COLORES = {  # fondo, texto
    "M": ("FFF2CC", "7F6000"),
    "T": ("DDEBF7", "1F4E78"),
    "N": ("3A3A5C", "FFFFFF"),
    "L": ("F2F2F2", "808080"),
    "F": ("C6EFCE", "006100"),
}
NOMBRES = {
    "M": "Mañana",
    "T": "Tarde",
    "N": "Noche",
    "L": "Libranza de la rotación",
    "F": "Fiesta / libre de ajuste de horas",
}

AZUL = "1F4E78"
ROJO = "C00000"
fill = lambda color: PatternFill("solid", start_color=color, end_color=color)  # noqa: E731
FILL_CABECERA = fill(AZUL)
FILL_FESTIVO = fill("F8CBAD")
FILL_FINDE = fill("D9D9D9")
FILL_ENTRADA = fill("FFFDE7")  # casillas que rellena el usuario
FILL_ERROR = fill("FFC7CE")
FILL_OK = fill("C6EFCE")
FONT_CABECERA = Font(bold=True, color="FFFFFF")
FONT_TITULO = Font(bold=True, size=16, color=AZUL)
FONT_NOTA = Font(italic=True, color="595959")
CENTRO = Alignment(horizontal="center", vertical="center")
AJUSTAR = Alignment(wrap_text=True, vertical="top")
fino = Side(style="thin", color="BFBFBF")
grueso = Side(style="medium", color="404040")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)


def colorear_turnos(ws, rango: str, codigos=("M", "T", "N", "L", "F")):
    """Formato condicional: el color sigue a la letra aunque se cambie a mano."""
    for c in codigos:
        fondo, texto = COLORES[c]
        ws.conditional_formatting.add(
            rango,
            CellIsRule(operator="equal", formula=[f'"{c}"'], fill=fill(fondo),
                       font=Font(color=texto, bold=c == "N")),
        )


def desplegable(ws, rango: str, opciones: list[str], titulo: str, mensaje: str):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True)
    dv.error, dv.errorTitle = f"Escribe una de estas opciones: {', '.join(opciones)}", "Valor no válido"
    dv.prompt, dv.promptTitle = mensaje, titulo
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rango)


def cabecera(ws, fila: int, textos: list[str], col_ini: int = 1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col_ini + i, t)
        c.fill, c.font, c.alignment, c.border = FILL_CABECERA, FONT_CABECERA, Alignment(
            horizontal="center", vertical="center", wrap_text=True), BORDE

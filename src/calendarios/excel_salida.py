"""Escribe el calendario final (calendario_AAAA.xlsx).

Todo lo que se puede recalcular va con fórmulas: si se cambia un turno a mano,
las horas, los mínimos, los festivos y los colores se actualizan solos.
"""

from __future__ import annotations

import datetime as dt
import re

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from . import estilos as E
from .datos import Datos
from .modelo import Resultado

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre",
         "Octubre", "Noviembre", "Diciembre"]
LETRA_DIA = "LMXJVSD"

# Disposición de las hojas de grupo
FILA_MES, FILA_FECHA, FILA_SEMANA, FILA_FESTIVO = 3, 4, 5, 6
FILA_PRIMERA_PERSONA = 7
COL_PRIMER_DIA = 10  # J
COLUMNAS_RESUMEN = ["Nº", "Nombre", "Horas", "Diferencia", "Mañanas", "Tardes", "Noches", "F",
                    "Festivos trabajados"]
MARCA_FESTIVO = "★"


def nombre_hoja(nombre: str) -> str:
    return re.sub(r"[\[\]:*?/\\]", "-", nombre)[:31]


def escribir_calendario(destino, datos: Datos, resultados: list[Resultado]):
    """`destino` puede ser una ruta o un fichero en memoria (BytesIO)."""
    wb = Workbook()
    resumen = wb.active
    hojas = []
    for r in resultados:
        ws = wb.create_sheet(nombre_hoja(r.grupo.nombre))
        _hoja_grupo(ws, datos, r)
        hojas.append(ws.title)
    _hoja_resumen(resumen, datos, resultados, hojas)
    _hoja_navidad(wb.create_sheet("Navidad"), datos, resultados, hojas)
    avisos = [(r.grupo.nombre, a) for r in resultados for a in r.avisos]
    if avisos:
        ws = wb.create_sheet("Avisos")
        E.cabecera(ws, 1, ["Grupo", "Aviso"])
        for i, (g, a) in enumerate(avisos, start=2):
            ws.cell(i, 1, g)
            ws.cell(i, 2, a)
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 120
    wb.save(destino)
    return destino


def _rango_dias(fila: int, n_dias: int, absoluto_col: bool = True) -> str:
    a = get_column_letter(COL_PRIMER_DIA)
    b = get_column_letter(COL_PRIMER_DIA + n_dias - 1)
    d = "$" if absoluto_col else ""
    return f"{d}{a}{fila}:{d}{b}{fila}"


def _leyenda(ws, fila: int, col: int):
    for i, codigo in enumerate(("M", "T", "N", "L", "F")):
        c = ws.cell(fila, col + i * 4, codigo)
        fondo, texto = E.COLORES[codigo]
        c.fill, c.font, c.alignment, c.border = E.fill(fondo), Font(bold=True, color=texto), E.CENTRO, E.BORDE
        ws.cell(fila, col + i * 4 + 1, E.NOMBRES[codigo]).font = Font(size=9)


def _hoja_grupo(ws, datos: Datos, r: Resultado):
    p = datos.parametros
    n = len(r.fechas)
    personas = r.grupo.personas
    ult = FILA_PRIMERA_PERSONA + len(personas) - 1
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90

    ws["A1"] = f"Calendario {p.anio} · {r.grupo.nombre}"
    ws["A1"].font = E.FONT_TITULO
    _leyenda(ws, 2, COL_PRIMER_DIA)
    ws.cell(2, COL_PRIMER_DIA + 21, "Letra subrayada = cambio respecto a la rotación.   "
                                    f"{MARCA_FESTIVO} = festivo").font = E.FONT_NOTA

    # Cabeceras fijas (A..I), combinadas en vertical
    for i, texto in enumerate(COLUMNAS_RESUMEN, start=1):
        ws.merge_cells(start_row=FILA_MES, start_column=i, end_row=FILA_FESTIVO, end_column=i)
        c = ws.cell(FILA_MES, i, texto)
        c.fill, c.font, c.border = E.FILL_CABECERA, E.FONT_CABECERA, E.BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Cabeceras de días
    inicio_mes = COL_PRIMER_DIA
    for d, fecha in enumerate(r.fechas):
        col = COL_PRIMER_DIA + d
        festivo = datos.festivos.get(fecha)
        finde = fecha.weekday() >= 5
        c_fecha = ws.cell(FILA_FECHA, col, fecha)
        c_fecha.number_format = "d"
        c_sem = ws.cell(FILA_SEMANA, col, LETRA_DIA[fecha.weekday()])
        c_fest = ws.cell(FILA_FESTIVO, col, MARCA_FESTIVO if festivo else None)
        for c in (c_fecha, c_sem, c_fest):
            c.alignment, c.border = E.CENTRO, E.BORDE
            c.font = Font(bold=True, size=9, color=E.ROJO if festivo else "000000")
            if festivo:
                c.fill = E.FILL_FESTIVO
            elif finde:
                c.fill = E.FILL_FINDE
        if festivo:
            c_fest.comment = Comment(festivo, "calendarios")
        ultimo_del_mes = d == n - 1 or r.fechas[d + 1].month != fecha.month
        if ultimo_del_mes:
            ws.merge_cells(start_row=FILA_MES, start_column=inicio_mes, end_row=FILA_MES, end_column=col)
            c = ws.cell(FILA_MES, inicio_mes, f"{MESES[fecha.month - 1]} {fecha.year}")
            c.fill, c.font, c.alignment = E.FILL_CABECERA, E.FONT_CABECERA, E.CENTRO
            inicio_mes = col + 1
        ws.column_dimensions[get_column_letter(col)].width = 3.4

    # Turnos
    for k, pe in enumerate(personas):
        f = FILA_PRIMERA_PERSONA + k
        ws.cell(f, 1, pe.numero).alignment = E.CENTRO
        ws.cell(f, 2, pe.nombre).font = Font(bold=True)
        for d in range(n):
            s = r.turnos[k][d]
            c = ws.cell(f, COL_PRIMER_DIA + d, s)
            c.alignment = E.CENTRO
            primero_mes = r.fechas[d].day == 1
            c.border = Border(left=E.grueso if primero_mes else E.fino, right=E.fino, top=E.fino, bottom=E.fino)
            if s in ("M", "T", "N") and s != r.base[k][d]:
                c.font = Font(underline="double", bold=True)
        _formulas_persona(ws, f, n, p)
    rango_turnos = f"{get_column_letter(COL_PRIMER_DIA)}{FILA_PRIMERA_PERSONA}:" \
                   f"{get_column_letter(COL_PRIMER_DIA + n - 1)}{ult}"
    E.colorear_turnos(ws, rango_turnos)
    E.desplegable(ws, rango_turnos, ["M", "T", "N", "L", "F"], "Turno",
                  "M mañana · T tarde · N noche · L libre · F fiesta")

    # Cobertura diaria
    fila_cob = ult + 2
    ws.cell(fila_cob - 1, 2, "Personas por turno (en rojo si no llega al mínimo)").font = E.FONT_NOTA
    for i, (s, nombre) in enumerate((("M", "Mañanas"), ("T", "Tardes"), ("N", "Noches"))):
        f = fila_cob + i
        ws.cell(f, 2, nombre).font = Font(bold=True)
        ws.cell(f, 3, f"mín. {p.minimos[s]}").font = E.FONT_NOTA
        for d in range(n):
            col = get_column_letter(COL_PRIMER_DIA + d)
            c = ws.cell(f, COL_PRIMER_DIA + d, f'=COUNTIF({col}${FILA_PRIMERA_PERSONA}:{col}${ult},"{s}")')
            c.alignment, c.border, c.font = E.CENTRO, E.BORDE, Font(size=9)
        rango = _rango_dias(f, n, absoluto_col=False)
        primera = f"{get_column_letter(COL_PRIMER_DIA)}{f}"
        ws.conditional_formatting.add(rango, FormulaRule(formula=[f"{primera}<{p.minimos[s]}"],
                                                         fill=E.FILL_ERROR, font=Font(bold=True, color=E.ROJO)))
    f = fila_cob + 3
    ws.cell(f, 2, "Días por debajo del mínimo").font = Font(bold=True)
    c = ws.cell(f, 5, "=" + "+".join(
        f'COUNTIF({_rango_dias(fila_cob + i, n)},"<{p.minimos[s]}")' for i, s in enumerate("MTN")))
    c.font = Font(bold=True)
    ws.conditional_formatting.add(f"E{f}", FormulaRule(formula=[f"E{f}>0"], fill=E.FILL_ERROR))
    ws.conditional_formatting.add(f"E{f}", FormulaRule(formula=[f"E{f}=0"], fill=E.FILL_OK))

    # Diferencia de horas en rojo si se sale del margen
    ws.conditional_formatting.add(
        f"D{FILA_PRIMERA_PERSONA}:D{ult}",
        FormulaRule(formula=[f"ABS(ROUND(C{FILA_PRIMERA_PERSONA}*1440,0)-{p.horas_anuales})>{p.margen}"],
                    fill=E.FILL_ERROR, font=Font(bold=True, color=E.ROJO)))
    ws.conditional_formatting.add(
        f"D{FILA_PRIMERA_PERSONA}:D{ult}",
        FormulaRule(formula=[f"ABS(ROUND(C{FILA_PRIMERA_PERSONA}*1440,0)-{p.horas_anuales})<={p.margen}"],
                    fill=E.FILL_OK))

    for col, w in zip("ABCDEFGHI", (4, 22, 8, 10, 8, 7, 7, 5, 9)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[FILA_MES].height = 20
    ws.freeze_panes = ws.cell(FILA_PRIMERA_PERSONA, COL_PRIMER_DIA)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_cols = "A:B"


def _formulas_persona(ws, f: int, n: int, p) -> None:
    fila = _rango_dias(f, n)
    cnt = lambda s: f'COUNTIF({fila},"{s}")'  # noqa: E731
    minutos = f"({cnt('M')}*{p.duracion['M']}+{cnt('T')}*{p.duracion['T']}+{cnt('N')}*{p.duracion['N']})"
    c = ws.cell(f, 3, f"={minutos}/1440")
    c.number_format = "[h]:mm"
    dif = f"(ROUND(C{f}*1440,0)-{p.horas_anuales})"
    # sin TEXT(): sus códigos de formato dependen del idioma de Excel
    ws.cell(f, 4, f'=IF({dif}<0,"-","+")&INT(ABS({dif})/60)&":"&RIGHT("0"&MOD(ABS({dif}),60),2)')
    for col, s in zip((5, 6, 7, 8), "MTNF"):
        ws.cell(f, col, f"={cnt(s)}")
    festivos = _rango_dias(FILA_FESTIVO, n)
    ws.cell(f, 9, f'=SUMPRODUCT(({festivos}="{MARCA_FESTIVO}")*(({fila}="M")+({fila}="T")+({fila}="N")))')
    for col in range(1, 10):
        c = ws.cell(f, col)
        c.border = E.BORDE
        if col != 2:
            c.alignment = E.CENTRO


def _col_fecha(r: Resultado, fecha: dt.date) -> str:
    return get_column_letter(COL_PRIMER_DIA + r.fechas.index(fecha))


def _hoja_resumen(ws, datos: Datos, resultados: list[Resultado], hojas: list[str]):
    anio = datos.parametros.anio
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Calendario {anio} · Resumen"
    ws["A1"].font = E.FONT_TITULO
    h, mg = datos.parametros.horas_anuales, datos.parametros.margen
    ws["A2"] = (f"Objetivo: {h // 60}:{h % 60:02d} h por persona (± {mg // 60}:{mg % 60:02d}). "
                "Se actualiza solo si cambias turnos en las hojas de cada grupo.")
    ws["A2"].font = E.FONT_NOTA
    _leyenda(ws, 3, 1)
    cab = ["Grupo", "Nº", "Nombre", "Horas", "Diferencia", "Mañanas", "Tardes", "Noches", "F",
           "Festivos trabajados", "Nochebuena", "Navidad", "Nochevieja", "Días bajo mínimos (grupo)"]
    E.cabecera(ws, 5, cab)
    ws.row_dimensions[5].height = 32
    f = 6
    for r, hoja in zip(resultados, hojas):
        ref = f"'{hoja}'!"
        fila_bajo_min = FILA_PRIMERA_PERSONA + len(r.grupo.personas) + 4
        for k, pe in enumerate(r.grupo.personas):
            fg = FILA_PRIMERA_PERSONA + k
            valores = [r.grupo.nombre, f"={ref}A{fg}", f"={ref}B{fg}", f"={ref}C{fg}", f"={ref}D{fg}",
                       f"={ref}E{fg}", f"={ref}F{fg}", f"={ref}G{fg}", f"={ref}H{fg}", f"={ref}I{fg}"]
            for dia in (24, 25, 31):
                valores.append(f"={ref}{_col_fecha(r, dt.date(anio, 12, dia))}{fg}")
            valores.append(f"={ref}E{fila_bajo_min}" if k == 0 else None)
            for col, v in enumerate(valores, start=1):
                c = ws.cell(f, col, v)
                c.border = E.BORDE
                if col not in (1, 3):
                    c.alignment = E.CENTRO
            ws.cell(f, 4).number_format = "[h]:mm"
            if k == 0:
                for col in range(1, len(cab) + 1):
                    ws.cell(f, col).border = Border(top=E.grueso, left=E.fino, right=E.fino, bottom=E.fino)
            f += 1
    ult = f - 1
    E.colorear_turnos(ws, f"K6:M{ult}", ("M", "T", "N", "L", "F"))
    ws.conditional_formatting.add(f"N6:N{ult}", FormulaRule(formula=["AND(N6<>\"\",N6>0)"], fill=E.FILL_ERROR))
    ws.conditional_formatting.add(f"N6:N{ult}", FormulaRule(formula=["AND(N6<>\"\",N6=0)"], fill=E.FILL_OK))
    ws.conditional_formatting.add(
        f"E6:E{ult}", FormulaRule(formula=[f"ABS(ROUND(D6*1440,0)-{h})>{mg}"], fill=E.FILL_ERROR))
    ws.conditional_formatting.add(
        f"E6:E{ult}", FormulaRule(formula=[f"ABS(ROUND(D6*1440,0)-{h})<={mg}"], fill=E.FILL_OK))
    for col, w in zip("ABCDEFGHIJKLMN", (14, 5, 24, 9, 11, 9, 8, 8, 6, 11, 12, 10, 12, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D6"


def _hoja_navidad(ws, datos: Datos, resultados: list[Resultado], hojas: list[str]):
    anio = datos.parametros.anio
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Navidad {anio} y qué le toca a cada una en {anio + 1}"
    ws["A1"].font = E.FONT_TITULO
    ws["A2"] = ("Las que trabajan Nochebuena y Nochevieja de mañana libran Reyes el año siguiente y quedan fuera "
                "de la alternancia. Se actualiza si cambias turnos a mano.")
    ws["A2"].font = E.FONT_NOTA
    E.cabecera(ws, 4, ["Grupo", "Nombre", f"Nochebuena {anio}", f"Navidad {anio}", f"Nochevieja {anio}",
                       f"Reyes {anio}", f"En {anio + 1}"])
    f = 5
    trabaja = lambda c: f'OR({c}="M",{c}="T",{c}="N")'  # noqa: E731
    for r, hoja in zip(resultados, hojas):
        ref = f"'{hoja}'!"
        for k, pe in enumerate(r.grupo.personas):
            fg = FILA_PRIMERA_PERSONA + k
            ws.cell(f, 1, r.grupo.nombre)
            ws.cell(f, 2, pe.nombre)
            for col, fecha in ((3, dt.date(anio, 12, 24)), (4, dt.date(anio, 12, 25)),
                               (5, dt.date(anio, 12, 31)), (6, dt.date(anio, 1, 6))):
                ws.cell(f, col, f"={ref}{_col_fecha(r, fecha)}{fg}").alignment = E.CENTRO
            ws.cell(f, 7, (f'=IF(AND(C{f}="M",E{f}="M"),"Libra Reyes; fuera de la alternancia",'
                           f'IF(AND({trabaja(f"C{f}")},NOT({trabaja(f"E{f}")})),"Nochevieja",'
                           f'IF(AND({trabaja(f"E{f}")},NOT({trabaja(f"C{f}")})),"Nochebuena","Revisar")))'))
            for col in range(1, 8):
                ws.cell(f, col).border = E.BORDE
            f += 1
    E.colorear_turnos(ws, f"C5:F{f - 1}")
    for col, w in zip("ABCDEFG", (14, 24, 15, 13, 15, 11, 36)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

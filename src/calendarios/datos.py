"""Modelo de datos, valores por defecto y lectura del Excel de entrada (datos_AAAA.xlsx).

Los datos viajan como un diccionario JSON (lo usa la web) con esta forma:

    {"anio": 2027,
     "parametros": {"horas_anuales": 102900, "margen": 210, "duracion": {"M": 420, ...},   # minutos
                    "minimos": {"M": 3, "T": 3, "N": 1}, "descansos_noche": [3, 3, 3, 3, 2, 2, 3], ...},
     "festivos": [{"fecha": "2027-01-01", "nombre": "Año Nuevo", "tipo": "Nacional"}, ...],
     "grupos": [{"nombre": "3ª-4ª C", "semana_inicio": 18, "salto": 2,
                 "rotacion": [["M", "M", ...7], ...26],          # "" = casilla vacía
                 "personas": [{"nombre": "...", "nochebuena": "M", "nochevieja": "", "mananas": false}]}]}
"""

from __future__ import annotations

import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

TURNOS_TRABAJO = ("M", "T", "N")
CODIGOS_ROTACION = ("M", "T", "N", "L")
SEMANAS_ROTACION = 26
NUM_GRUPOS = 5
PERSONAS_POR_GRUPO = 13
SEMANA_INICIO_POR_DEFECTO = 18
DIAS_SEMANA = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]

HOJA_PARAMETROS = "Parámetros"
HOJA_FESTIVOS = "Festivos"
HOJA_GRUPOS = "Grupos"
HOJA_PERSONAS = "Personas"
FILA_PRIMERA_SEMANA = 4  # en las hojas "Rotación N", la semana 1 está en la fila 4 (columnas B-H)

# Rotación "3ª-4ª C" (transcrita de la hoja en papel)
ROTACION_3_4_C = (
    "MMMMMTT LLMMTNN LLMTTLL MMMLLTT LLMMMLL LLTTTLL MMMMMTT LLMMMLL NNLLMMM TTTLLMM MMTTNLL TTLLMMM TTNNLLL "
    "MMMMMTT LLMMTNN LLMTTLL MMMLLTT LLMMMMM LLTTTLL MMMMMTT LLMMMLL NNLLMMM TTTLLMM MMTTNLL TTLLMMM TTNNLLL"
).split()


def hoja_rotacion(numero: int) -> str:
    return f"Rotación {numero}"


class ErrorDatos(Exception):
    """Error en los datos de entrada, con un mensaje pensado para quien los rellena."""


@dataclass
class Parametros:
    anio: int = 2027
    horas_anuales: int = 1715 * 60  # todo en minutos
    margen: int = 3 * 60 + 30
    duracion: dict[str, int] = field(default_factory=lambda: {"M": 420, "T": 420, "N": 682})
    minimos: dict[str, int] = field(default_factory=lambda: {"M": 3, "T": 3, "N": 1})
    # Días de descanso tras la última noche según el día de esa noche (lunes..domingo).
    # Si en la rotación ese día había turno, sale como F.
    descansos_noche: list[int] = field(default_factory=lambda: [3, 3, 3, 3, 2, 2, 3])
    max_dias_seguidos: int = 7
    libranzas_tras_max: int = 2
    tiempo_max_s: int = 60


@dataclass
class Persona:
    numero: int
    nombre: str
    nochebuena_prev: str | None = None  # turno (M/T/N) en Nochebuena del año anterior
    nochevieja_prev: str | None = None
    mananas_navidad: bool = False  # elegida para trabajar 24 y 31 de mañana este año

    @property
    def mananas_navidad_prev(self) -> bool:
        return self.nochebuena_prev == "M" and self.nochevieja_prev == "M"


@dataclass
class Grupo:
    numero: int
    nombre: str
    fila_inicial: int  # semana (1-26) de la rotación de la persona 1 en la primera semana del año
    salto: int
    rotacion: list[str]  # 26 cadenas de 7 letras (lunes..domingo)
    personas: list[Persona]


@dataclass
class Datos:
    parametros: Parametros
    festivos: dict[dt.date, str]
    grupos: list[Grupo]

    @property
    def lunes_referencia(self) -> dt.date:
        return lunes_de(dt.date(self.parametros.anio, 1, 1))


def lunes_de(fecha: dt.date) -> dt.date:
    return fecha - dt.timedelta(days=fecha.weekday())


def hm(minutos: int) -> str:
    return f"{minutos // 60}:{minutos % 60:02d}"


# ---------------------------------------------------------------- festivos

def domingo_de_pascua(anio: int) -> dt.date:
    a, b, c = anio % 19, anio // 100, anio % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = (h + l - 7 * m + 114) % 31 + 1
    return dt.date(anio, mes, dia)


def festivos_por_defecto(anio: int) -> list[tuple[dt.date, str, str]]:
    """Festivos nacionales + locales de Burgos. Hay que revisarlos cada año con el BOCYL."""
    pascua = domingo_de_pascua(anio)
    d = lambda mes, dia: dt.date(anio, mes, dia)  # noqa: E731
    lista = [
        (d(1, 1), "Año Nuevo", "Nacional"),
        (d(1, 6), "Reyes", "Nacional"),
        (pascua - dt.timedelta(days=3), "Jueves Santo", "Nacional"),
        (pascua - dt.timedelta(days=2), "Viernes Santo", "Nacional"),
        (d(5, 1), "Fiesta del Trabajo", "Nacional"),
        # Curpillos: viernes siguiente al domingo del Corpus
        (pascua + dt.timedelta(days=68), "Curpillos", "Local"),
        (d(6, 29), "San Pedro y San Pablo", "Local"),
        (d(8, 15), "Asunción", "Nacional"),
        (d(10, 12), "Fiesta Nacional", "Nacional"),
        (d(11, 1), "Todos los Santos", "Nacional"),
        (d(12, 6), "Constitución", "Nacional"),
        (d(12, 8), "Inmaculada", "Nacional"),
        (d(12, 24), "Nochebuena", "Local"),
        (d(12, 25), "Navidad", "Nacional"),
        (d(12, 31), "Nochevieja", "Local"),
    ]
    return sorted(lista)


# ---------------------------------------------------------------- diccionario (JSON)

def parametros_a_dict(p: Parametros) -> dict:
    return {
        "horas_anuales": p.horas_anuales, "margen": p.margen, "duracion": dict(p.duracion),
        "minimos": dict(p.minimos), "descansos_noche": list(p.descansos_noche),
        "max_dias_seguidos": p.max_dias_seguidos, "libranzas_tras_max": p.libranzas_tras_max,
        "tiempo_max_s": p.tiempo_max_s,
    }


def festivos_dict(anio: int) -> list[dict]:
    return [{"fecha": f.isoformat(), "nombre": n, "tipo": t} for f, n, t in festivos_por_defecto(anio)]


def dict_por_defecto(anio: int, vacia: bool = False) -> dict:
    grupos = []
    for g in range(1, NUM_GRUPOS + 1):
        con_rot = g == 1 and not vacia
        grupos.append({
            "nombre": "3ª-4ª C" if con_rot else f"Grupo {g}",
            "semana_inicio": SEMANA_INICIO_POR_DEFECTO,
            "salto": 2,
            "rotacion": [list(s) if con_rot else [""] * 7 for s in ROTACION_3_4_C],
            "personas": [{"nombre": f"Persona {i}", "nochebuena": "", "nochevieja": "", "mananas": False}
                         for i in range(1, PERSONAS_POR_GRUPO + 1)],
        })
    return {"anio": anio, "parametros": parametros_a_dict(Parametros(anio=anio)),
            "festivos": festivos_dict(anio), "grupos": grupos}


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _turno(valor, donde: str) -> str | None:
    t = _texto(valor).upper()
    if t in ("", "-", "—", "NO", "L"):
        return None
    if t not in TURNOS_TRABAJO:
        raise ErrorDatos(f"{donde}: el turno debe ser M, T o N (o vacío), pero pone «{valor}»")
    return t


def _fecha(valor, donde: str) -> dt.date:
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    texto = _texto(valor)
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(texto, formato).date()
        except ValueError:
            pass
    raise ErrorDatos(f"{donde}: «{valor}» no es una fecha")


def datos_desde_dict(d: dict) -> Datos:
    """Valida el diccionario y lo convierte en Datos. Los grupos con la rotación vacía se ignoran."""
    try:
        anio = int(d["anio"])
    except (KeyError, TypeError, ValueError):
        raise ErrorDatos("Falta el año") from None
    q = d.get("parametros") or {}
    base = Parametros(anio=anio)
    p = Parametros(
        anio=anio,
        horas_anuales=int(q.get("horas_anuales", base.horas_anuales)),
        margen=int(q.get("margen", base.margen)),
        duracion={s: int((q.get("duracion") or {}).get(s, base.duracion[s])) for s in TURNOS_TRABAJO},
        minimos={s: int((q.get("minimos") or {}).get(s, base.minimos[s])) for s in TURNOS_TRABAJO},
        descansos_noche=[int(x) for x in q.get("descansos_noche", base.descansos_noche)],
        max_dias_seguidos=int(q.get("max_dias_seguidos", base.max_dias_seguidos)),
        libranzas_tras_max=int(q.get("libranzas_tras_max", base.libranzas_tras_max)),
        tiempo_max_s=int(q.get("tiempo_max_s", base.tiempo_max_s)),
    )
    if len(p.descansos_noche) != 7:
        raise ErrorDatos("Los descansos tras la noche deben tener un valor para cada día de la semana")

    festivos: dict[dt.date, str] = {}
    for i, f in enumerate(d.get("festivos") or [], start=1):
        if not _texto(f.get("fecha")):
            continue
        fecha = _fecha(f["fecha"], f"Festivo {i}")
        if fecha.year != anio:
            raise ErrorDatos(f"El festivo {fecha:%d/%m/%Y} ({f.get('nombre') or ''}) no es del año {anio}")
        festivos[fecha] = _texto(f.get("nombre")) or "Festivo"

    grupos = []
    for num, g in enumerate(d.get("grupos") or [], start=1):
        nombre = _texto(g.get("nombre")) or f"Grupo {num}"
        filas = g.get("rotacion") or []
        celdas = []
        for s in range(SEMANAS_ROTACION):
            fila = filas[s] if s < len(filas) else []
            celdas += [_texto(fila[dia]).upper() if dia < len(fila) else "" for dia in range(7)]
        if not any(celdas):
            continue  # grupo sin rotación todavía
        vacias = [f"semana {i // 7 + 1} {DIAS_SEMANA[i % 7]}" for i, c in enumerate(celdas) if not c]
        if vacias:
            raise ErrorDatos(f"{nombre}: la rotación tiene casillas vacías ({', '.join(vacias[:3])}"
                             f"{'…' if len(vacias) > 3 else ''})")
        malas = [c for c in celdas if c not in CODIGOS_ROTACION]
        if malas:
            raise ErrorDatos(f"{nombre}: en la rotación solo vale M, T, N o L (hay «{malas[0]}»)")
        rotacion = ["".join(celdas[s * 7:(s + 1) * 7]) for s in range(SEMANAS_ROTACION)]
        try:
            semana = int(g.get("semana_inicio") or 0)
        except (TypeError, ValueError):
            semana = 0
        if not 1 <= semana <= SEMANAS_ROTACION:
            raise ErrorDatos(f"{nombre}: la semana de arranque debe ser un número del 1 al 26")
        personas = []
        for i, pe in enumerate(g.get("personas") or [], start=1):
            if not _texto(pe.get("nombre")):
                continue
            donde = f"{nombre}, {_texto(pe.get('nombre'))}"
            personas.append(Persona(
                numero=i, nombre=_texto(pe["nombre"]),
                nochebuena_prev=_turno(pe.get("nochebuena"), donde + " (Nochebuena)"),
                nochevieja_prev=_turno(pe.get("nochevieja"), donde + " (Nochevieja)"),
                mananas_navidad=bool(pe.get("mananas")),
            ))
        if not personas:
            raise ErrorDatos(f"{nombre}: tiene rotación pero no tiene personas")
        nombres = [x.nombre for x in personas]
        repetidos = {x for x in nombres if nombres.count(x) > 1}
        if repetidos:
            raise ErrorDatos(f"{nombre}: hay nombres repetidos ({', '.join(sorted(repetidos))})")
        if sum(x.mananas_navidad for x in personas) > 2:
            raise ErrorDatos(f"{nombre}: hay más de 2 personas marcadas para las mañanas de Navidad")
        grupos.append(Grupo(num, nombre, semana, int(g.get("salto") or 2), rotacion, personas))

    if not grupos:
        raise ErrorDatos("No hay ningún grupo con la rotación completa")
    return Datos(p, festivos, grupos)


# ---------------------------------------------------------------- Excel -> diccionario

# Etiquetas de la hoja "Parámetros" (columna A)
ETIQUETAS_PARAMETROS = [
    "Año", "Horas anuales", "Margen permitido (+/-)",
    "Duración turno de mañana", "Duración turno de tarde", "Duración turno de noche",
    "Mínimo de personas de mañana", "Mínimo de personas de tarde", "Mínimo de personas de noche",
    *[f"Descansos si la última noche es en {d}" for d in DIAS_SEMANA],
    "Máximo de días seguidos trabajando", "Libranzas después del máximo de días seguidos",
    "Tiempo máximo de cálculo por grupo (segundos)",
]


def a_minutos(valor, etiqueta: str) -> int:
    """Acepta '7:00', '11:22', 1715, 3.5, horas de Excel (time/timedelta/fecha)..."""
    if isinstance(valor, dt.timedelta):
        return round(valor.total_seconds() / 60)
    if isinstance(valor, dt.datetime):  # Excel guarda >24h como fecha desde 1899-12-30
        return round((valor - dt.datetime(1899, 12, 30)).total_seconds() / 60)
    if isinstance(valor, dt.time):
        return valor.hour * 60 + valor.minute
    if isinstance(valor, (int, float)):
        return round(valor * 60)
    if isinstance(valor, str):
        texto = valor.strip().lower().replace(",", ".")
        m = re.fullmatch(r"(\d+)\s*(?::|h)\s*(\d{1,2})?\s*(?:min)?", texto)
        if m:
            return int(m.group(1)) * 60 + int(m.group(2) or 0)
        try:
            return round(float(texto) * 60)
        except ValueError:
            pass
    raise ErrorDatos(f"'{etiqueta}': no entiendo el valor «{valor}». Escríbelo como horas:minutos, p. ej. 11:22")


def abrir_excel(origen):
    from zipfile import BadZipFile

    from openpyxl import load_workbook

    try:
        return load_workbook(origen, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        raise ErrorDatos("Ese fichero no es un Excel (.xlsx) válido. Si lo tienes en otro formato, "
                         "ábrelo con Excel y guárdalo como «Libro de Excel (.xlsx)».") from None


def dict_desde_excel(origen: str | Path | bytes) -> dict:
    if isinstance(origen, (bytes, bytearray)):
        nombre_fichero = "el Excel"
        origen = io.BytesIO(origen)
    else:
        origen = Path(origen)
        if not origen.exists():
            raise ErrorDatos(f"No encuentro el fichero {origen}")
        nombre_fichero = origen.name
    wb = abrir_excel(origen)
    for hoja in (HOJA_PARAMETROS, HOJA_FESTIVOS, HOJA_GRUPOS, HOJA_PERSONAS):
        if hoja not in wb.sheetnames:
            raise ErrorDatos(f"Falta la hoja «{hoja}» en {nombre_fichero}. ¿Es la plantilla de datos?")

    valores = {}
    for fila in wb[HOJA_PARAMETROS].iter_rows(min_row=2, max_col=2, values_only=True):
        if _texto(fila[0]) in ETIQUETAS_PARAMETROS and _texto(fila[1]) != "":
            valores[_texto(fila[0])] = fila[1]
    if "Año" not in valores:
        raise ErrorDatos("Parámetros: falta el año")
    anio = int(valores["Año"])
    p = parametros_a_dict(Parametros(anio=anio))
    tiempos = {"Horas anuales": ("horas_anuales", None), "Margen permitido (+/-)": ("margen", None),
               "Duración turno de mañana": ("duracion", "M"), "Duración turno de tarde": ("duracion", "T"),
               "Duración turno de noche": ("duracion", "N")}
    for etiqueta, v in valores.items():
        if etiqueta in tiempos:
            clave, sub = tiempos[etiqueta]
            minutos = a_minutos(v, etiqueta)
            if sub:
                p[clave][sub] = minutos
            else:
                p[clave] = minutos
        elif etiqueta.startswith("Mínimo de personas de "):
            p["minimos"]["MTN"[["mañana", "tarde", "noche"].index(etiqueta.split()[-1])]] = int(v)
        elif etiqueta.startswith("Descansos si la última noche es en "):
            p["descansos_noche"][DIAS_SEMANA.index(etiqueta.split()[-1])] = int(v)
        elif etiqueta == "Máximo de días seguidos trabajando":
            p["max_dias_seguidos"] = int(v)
        elif etiqueta == "Libranzas después del máximo de días seguidos":
            p["libranzas_tras_max"] = int(v)
        elif etiqueta.startswith("Tiempo máximo"):
            p["tiempo_max_s"] = int(v)

    festivos = []
    for i, fila in enumerate(wb[HOJA_FESTIVOS].iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        if fila[0] is None:
            continue
        festivos.append({"fecha": _fecha(fila[0], f"Festivos, fila {i}").isoformat(),
                         "nombre": _texto(fila[1]), "tipo": _texto(fila[2])})

    personas: dict[int, list[dict]] = {}
    for fila in wb[HOJA_PERSONAS].iter_rows(min_row=2, max_col=7, values_only=True):
        g, _, _, nombre, nb, nv, manana = fila
        if g is None:
            continue
        personas.setdefault(int(g), []).append({
            "nombre": _texto(nombre), "nochebuena": _texto(nb).upper(), "nochevieja": _texto(nv).upper(),
            "mananas": _texto(manana).upper() in ("SÍ", "SI", "S", "X"),
        })

    grupos = []
    for fila in wb[HOJA_GRUPOS].iter_rows(min_row=2, max_col=4, values_only=True):
        num, nombre, semana, salto = fila
        if num is None:
            continue
        num = int(num)
        rot = [[""] * 7 for _ in range(SEMANAS_ROTACION)]
        if hoja_rotacion(num) in wb.sheetnames:
            ws = wb[hoja_rotacion(num)]
            rot = [[_texto(ws.cell(FILA_PRIMERA_SEMANA + s, 2 + dia).value).upper() for dia in range(7)]
                   for s in range(SEMANAS_ROTACION)]
        grupos.append({"nombre": _texto(nombre) or f"Grupo {num}",
                       "semana_inicio": int(semana) if _texto(semana) else SEMANA_INICIO_POR_DEFECTO,
                       "salto": int(salto or 2), "rotacion": rot, "personas": personas.get(num, [])})
    return {"anio": anio, "parametros": p, "festivos": festivos, "grupos": grupos}


def leer_datos(ruta: str | Path) -> Datos:
    return datos_desde_dict(dict_desde_excel(ruta))

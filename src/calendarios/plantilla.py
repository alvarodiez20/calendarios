"""Crea el Excel de entrada (datos_AAAA.xlsx) listo para rellenar."""

from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import estilos as E
from .datos import (DIAS_SEMANA, ETIQUETAS_PARAMETROS, FILA_PRIMERA_SEMANA, HOJA_FESTIVOS, HOJA_GRUPOS,
                    HOJA_PARAMETROS, HOJA_PERSONAS, SEMANAS_ROTACION, _fecha, dict_por_defecto, hm,
                    hoja_rotacion, lunes_de)

DIAS = [d.capitalize() for d in DIAS_SEMANA]


def _fila_parametro(etiqueta: str) -> int:
    return ETIQUETAS_PARAMETROS.index(etiqueta) + 2


def _ref_param(etiqueta: str) -> str:
    return f"'{HOJA_PARAMETROS}'!$B${_fila_parametro(etiqueta)}"


def _horas(etiqueta: str) -> str:
    """Fórmula que convierte el texto '11:22' (o una hora de Excel) en horas decimales."""
    ref = _ref_param(etiqueta)
    return f"(IFERROR(TIMEVALUE({ref}),{ref})*24)"


def escribir_datos(destino, datos: dict | None = None, anio: int | None = None, vacia: bool = False):
    """Escribe el Excel de datos. `destino` puede ser una ruta o un fichero en memoria (BytesIO)."""
    d = datos or dict_por_defecto(anio or 2027, vacia=vacia)
    anio = int(d["anio"])
    lunes = lunes_de(dt.date(anio, 1, 1))
    wb = Workbook()
    _hoja_instrucciones(wb.active, anio, lunes)
    _hoja_parametros(wb.create_sheet(HOJA_PARAMETROS), anio, d["parametros"])
    _hoja_festivos(wb.create_sheet(HOJA_FESTIVOS), d["festivos"])
    _hoja_grupos(wb.create_sheet(HOJA_GRUPOS), d["grupos"], lunes)
    _hoja_personas(wb.create_sheet(HOJA_PERSONAS), d["grupos"], anio)
    for i, g in enumerate(d["grupos"], start=1):
        _hoja_rotacion(wb.create_sheet(hoja_rotacion(i)), i, g)
    wb.save(destino)
    return destino


def _hoja_instrucciones(ws, anio: int, lunes: dt.date):
    ws.title = "Cómo usar"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 110
    ws["B2"] = f"Datos para el calendario {anio}"
    ws["B2"].font = E.FONT_TITULO
    pasos = [
        ("Solo hay que escribir en las casillas amarillas. Todo lo demás se calcula solo.", True),
        ("", False),
        ("1. Grupos: el nombre de cada grupo y la semana de su rotación (1 a 26) en la que está la persona 1 "
         f"la semana del lunes {lunes:%d/%m/%Y}.", False),
        ("2. Rotación 1 … Rotación 5: la rotación de 26 semanas de cada grupo (M, T, N o L). "
         "Debajo de cada rotación se ve en rojo si algún día no llega a los mínimos. "
         "Un grupo con la rotación vacía no se calcula.", False),
        ("3. Personas: los nombres, en orden (la persona 1 empieza en la semana indicada en Grupos, "
         "la 2 va dos semanas por delante, etc.).", False),
        (f"   · Nochebuena / Nochevieja {anio - 1}: el turno (M, T o N) que hizo cada una. Si lo dejas vacío, "
         "a esa persona no se le aplica la alternancia.", False),
        ("   · Las dos que hicieron Nochebuena Y Nochevieja de mañana se detectan solas: libran Reyes y quedan "
         "fuera de la alternancia.", False),
        (f"   · Mañanas de Navidad {anio}: marca «Sí» a las 2 que trabajarán el 24 y el 31 de mañana. "
         "Si no marcas a nadie, las elige el programa.", False),
        (f"4. Festivos: revisa la lista de {anio}. Puedes añadir o borrar filas.", False),
        ("5. Parámetros: horas anuales, duración de los turnos, mínimos, descansos… Normalmente no se tocan.", False),
        ("", False),
        ("Después, sube este Excel en la web y pulsa «Generar calendario».", True),
    ]
    for i, (texto, negrita) in enumerate(pasos, start=4):
        c = ws.cell(i, 2, texto)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(bold=negrita, size=12)
    ws["B4"].fill = E.FILL_ENTRADA


def _hoja_parametros(ws, anio: int, p: dict):
    E.cabecera(ws, 1, ["Parámetro", "Valor", "Explicación"])
    valores = {
        "Año": (anio, "Año del calendario."),
        "Horas anuales": (p["horas_anuales"] // 60 if p["horas_anuales"] % 60 == 0 else hm(p["horas_anuales"]),
                          "Horas que debe sumar cada persona en el año."),
        "Margen permitido (+/-)": (hm(p["margen"]), "Cuánto se puede pasar o quedar corto (horas:minutos)."),
        "Duración turno de mañana": (hm(p["duracion"]["M"]), "horas:minutos"),
        "Duración turno de tarde": (hm(p["duracion"]["T"]), "horas:minutos"),
        "Duración turno de noche": (hm(p["duracion"]["N"]), "horas:minutos"),
        "Mínimo de personas de mañana": (p["minimos"]["M"], "Por grupo y día."),
        "Mínimo de personas de tarde": (p["minimos"]["T"], "Por grupo y día."),
        "Mínimo de personas de noche": (p["minimos"]["N"], "Por grupo y día."),
        "Máximo de días seguidos trabajando": (p["max_dias_seguidos"], ""),
        "Libranzas después del máximo de días seguidos": (p["libranzas_tras_max"], ""),
        "Mínimo de días libres seguidos": (p["min_libranzas_seguidas"],
                                           "Nunca se libra un día suelto (los festivos sí pueden librarse solos)."),
        "Máximo de días libres seguidos": (p["max_libranzas_seguidas"],
                                           "No se encadenan más días libres de los que ya da la rotación."),
        "Mínimo de días de trabajo seguidos": (p["min_dias_trabajo_seguidos"],
                                               "Nadie va a trabajar un día suelto entre dos libranzas."),
        "Tiempo máximo de cálculo por grupo (segundos)": (p["tiempo_max_s"], "Normalmente tarda unos segundos."),
    }
    for i, dia in enumerate(DIAS_SEMANA):
        valores[f"Descansos si la última noche es en {dia}"] = (
            p["descansos_noche"][i],
            "Días sin trabajar tras la última noche. Salen siempre como L, nunca como F."
            if i == 0 else "")
    for etiqueta in ETIQUETAS_PARAMETROS:
        fila = _fila_parametro(etiqueta)
        valor, nota = valores[etiqueta]
        ws.cell(fila, 1, etiqueta).border = E.BORDE
        c = ws.cell(fila, 2, valor)
        c.fill, c.border, c.alignment = E.FILL_ENTRADA, E.BORDE, E.CENTRO
        if isinstance(valor, str):
            c.number_format = "@"
        ws.cell(fila, 3, nota).font = E.FONT_NOTA
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 80


def _hoja_festivos(ws, festivos: list[dict]):
    E.cabecera(ws, 1, ["Fecha", "Festivo", "Tipo", "Día"])
    for i, fe in enumerate(festivos, start=2):
        ws.cell(i, 1, _fecha(fe["fecha"], "Festivo")).number_format = "dd/mm/yyyy"
        ws.cell(i, 2, fe.get("nombre"))
        ws.cell(i, 3, fe.get("tipo"))
    for i in range(2, 41):
        ws.cell(i, 1).number_format = "dd/mm/yyyy"
        ws.cell(i, 4, f'=IF(A{i}="","",CHOOSE(WEEKDAY(A{i},2),"lunes","martes","miércoles","jueves",'
                      f'"viernes","sábado","domingo"))').font = E.FONT_NOTA
        for col in (1, 2, 3):
            ws.cell(i, col).fill, ws.cell(i, col).border = E.FILL_ENTRADA, E.BORDE
    ws["F2"] = ("Revisa esta lista con el calendario laboral oficial (BOCYL). "
                "Si un festivo cae en domingo y se traslada, cambia la fecha.")
    ws["F2"].font = E.FONT_NOTA
    for col, w in zip("ABCD", (13, 26, 11, 11)):
        ws.column_dimensions[col].width = w


def _hoja_grupos(ws, grupos: list[dict], lunes: dt.date):
    E.cabecera(ws, 1, ["Nº", "Nombre del grupo",
                       f"Semana de arranque (semana de la rotación de la persona 1 el lunes {lunes:%d/%m/%Y})",
                       "Semanas de desfase entre personas", "Personas", "Rotación"])
    ws.row_dimensions[1].height = 60
    for i, g in enumerate(grupos, start=1):
        f = i + 1
        ws.cell(f, 1, i).alignment = E.CENTRO
        for col, valor in ((2, g["nombre"]), (3, g["semana_inicio"]), (4, g.get("salto", 2))):
            c = ws.cell(f, col, valor)
            c.fill, c.border, c.alignment = E.FILL_ENTRADA, E.BORDE, E.CENTRO
        ws.cell(f, 5, f'=COUNTIFS({HOJA_PERSONAS}!$A:$A,A{f},{HOJA_PERSONAS}!$D:$D,"?*")').alignment = E.CENTRO
        rng = f"'{hoja_rotacion(i)}'!$B${FILA_PRIMERA_SEMANA}:$H${FILA_PRIMERA_SEMANA + SEMANAS_ROTACION - 1}"
        ws.cell(f, 6, f'=IF(COUNTA({rng})=0,"Vacía (no se calcula)",IF(COUNTA({rng})=182,"Completa",'
                      f'"Faltan "&(182-COUNTA({rng}))&" casillas"))')
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(SEMANAS_ROTACION))
    dv.error = "Debe ser un número del 1 al 26"
    ws.add_data_validation(dv)
    dv.add(f"C2:C{len(grupos) + 1}")
    for col, w in zip("ABCDEF", (5, 24, 30, 16, 10, 22)):
        ws.column_dimensions[col].width = w


def _hoja_personas(ws, grupos: list[dict], anio: int):
    E.cabecera(ws, 1, ["Grupo nº", "Grupo", "Nº", "Nombre", f"Turno en Nochebuena {anio - 1}",
                       f"Turno en Nochevieja {anio - 1}", f"Mañanas de Navidad {anio}", f"Qué le toca en {anio}"])
    ws.row_dimensions[1].height = 45
    f = 2
    for gi, g in enumerate(grupos, start=1):
        for pi, pe in enumerate(g["personas"], start=1):
            ws.cell(f, 1, gi).alignment = E.CENTRO
            ws.cell(f, 2, f'=IFERROR(INDEX({HOJA_GRUPOS}!$B$2:$B$6,A{f}),"")').font = E.FONT_NOTA
            ws.cell(f, 3, pi).alignment = E.CENTRO
            for col, valor in ((4, pe.get("nombre")), (5, pe.get("nochebuena")), (6, pe.get("nochevieja")),
                               (7, "Sí" if pe.get("mananas") else "")):
                c = ws.cell(f, col, valor or None)
                c.fill, c.border = E.FILL_ENTRADA, E.BORDE
                if col > 4:
                    c.alignment = E.CENTRO
            ws.cell(f, 8, (
                f'=IF(AND(E{f}="M",F{f}="M"),"Libra Reyes; fuera de la alternancia",'
                f'IF(G{f}="Sí","24 y 31 de mañana",IF(AND(E{f}<>"",F{f}=""),"Nochevieja",'
                f'IF(AND(F{f}<>"",E{f}=""),"Nochebuena",IF(AND(E{f}<>"",F{f}<>""),"Revisar: hizo las dos",'
                f'"Lo decide el programa")))))')).font = E.FONT_NOTA
            if pi == 1:
                for col in range(1, 9):
                    ws.cell(f, col).border = E.Border(top=E.grueso, left=E.fino, right=E.fino, bottom=E.fino)
            f += 1
    ultima = max(f - 1, 2)
    E.desplegable(ws, f"E2:F{ultima}", ["M", "T", "N"], "Turno", "M, T o N. Vacío si no trabajó.")
    E.desplegable(ws, f"G2:G{ultima}", ["Sí"], "Mañanas de Navidad", "Marca «Sí» a 2 personas por grupo.")
    E.colorear_turnos(ws, f"E2:F{ultima}", ("M", "T", "N"))
    ws.freeze_panes = "E2"
    for col, w in zip("ABCDEFGH", (8, 14, 5, 28, 15, 15, 15, 34)):
        ws.column_dimensions[col].width = w


def _hoja_rotacion(ws, numero: int, g: dict):
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"Rotación del grupo {numero}:"
    ws["A1"].font = E.FONT_TITULO
    ws["D1"] = f"={HOJA_GRUPOS}!B{numero + 1}"
    ws["D1"].font = E.FONT_TITULO
    ws["A2"] = "Escribe M, T, N o L en cada casilla (26 semanas × 7 días)."
    ws["A2"].font = E.FONT_NOTA
    E.cabecera(ws, 3, ["Semana", *DIAS])
    f0, f1 = FILA_PRIMERA_SEMANA, FILA_PRIMERA_SEMANA + SEMANAS_ROTACION - 1
    rot = g.get("rotacion") or []
    for s in range(SEMANAS_ROTACION):
        f = f0 + s
        ws.cell(f, 1, s + 1).alignment = E.CENTRO
        ws.cell(f, 1).font = Font(bold=True)
        for d in range(7):
            valor = rot[s][d] if s < len(rot) and d < len(rot[s]) else ""
            c = ws.cell(f, 2 + d, valor or None)
            c.alignment, c.border = E.CENTRO, E.BORDE
    rango = f"B{f0}:H{f1}"
    E.colorear_turnos(ws, rango, ("M", "T", "N", "L"))
    E.desplegable(ws, rango, ["M", "T", "N", "L"], "Turno", "M mañana, T tarde, N noche, L libre")

    # Comprobaciones en vivo
    ws["J3"] = "Comprobaciones"
    ws["J3"].font = Font(bold=True, color=E.AZUL, size=12)
    cuenta = lambda s: f'COUNTIF({rango},"{s}")'  # noqa: E731
    horas_26 = (f"=({cuenta('M')}*{_horas('Duración turno de mañana')}+{cuenta('T')}*{_horas('Duración turno de tarde')}"
                f"+{cuenta('N')}*{_horas('Duración turno de noche')})")
    filas = [
        ("Horas por persona en 26 semanas", horas_26, "0.00"),
        ("Horas por persona en un año (≈ 2 vueltas)", "=K4*2", "0.00"),
        ("Horas que sobran → se quitan como días F", f"=K5-{_ref_param('Horas anuales')}", "0.00"),
        ("Días de mañana a quitar (aprox.)", f"=K6/{_horas('Duración turno de mañana')}", "0.0"),
    ]
    for i, (texto, formula, formato) in enumerate(filas, start=4):
        ws.cell(i, 10, texto)
        c = ws.cell(i, 11, formula)
        c.number_format, c.font = formato, Font(bold=True)

    fila = f1 + 3
    ws.cell(fila, 1, "Personas trabajando cada día (en rojo si no llega al mínimo)").font = Font(bold=True, color=E.AZUL)
    E.cabecera(ws, fila + 1, ["", *[d[:3] for d in DIAS]])
    f = fila + 2
    for paridad, texto in ((0, "Semanas impares"), (1, "Semanas pares")):
        for s, nombre, param in (("M", "Mañanas", "Mínimo de personas de mañana"),
                                 ("T", "Tardes", "Mínimo de personas de tarde"),
                                 ("N", "Noches", "Mínimo de personas de noche")):
            ws.cell(f, 1, f"{texto} · {nombre}")
            for d in range(7):
                col = get_column_letter(2 + d)
                r = f"{col}${f0}:{col}${f1}"
                c = ws.cell(f, 2 + d, f'=SUMPRODUCT((MOD(ROW({r})-{f0},2)={paridad})*({r}="{s}"))')
                c.alignment, c.border = E.CENTRO, E.BORDE
            ws.conditional_formatting.add(
                f"B{f}:H{f}", FormulaRule(formula=[f"B{f}<{_ref_param(param)}"], fill=E.FILL_ERROR))
            f += 1
    ws.cell(fila, 1).comment = Comment(
        "Con 13 personas desfasadas 2 semanas, cada semana trabajan las filas pares o las impares.", "calendarios")
    ws.column_dimensions["A"].width = 26
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 10
    ws.column_dimensions["J"].width = 42
    ws.column_dimensions["K"].width = 10
    ws.freeze_panes = "B4"
